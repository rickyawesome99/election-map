#!/usr/bin/env python3
"""
Louisiana's certified U.S. House results BY DISTRICT BY PARISH from the Secretary of State's
"human readable" results workbook (the static S3 export patch-county-house-2024-la-official.py
found; see its docstring). Tried for every House general and the December runoffs 2016–2024;
whichever workbooks exist are parsed. Writes data-entry/house_district_county_louisiana.csv,
one row per (year, district, parish). A runoff workbook's rows replace the November primary's
for the districts it decides, matching house_past_results.csv (which stores the decisive round).

Run from project root:  python3 scripts/fetch-house-district-county-louisiana.py
"""
import csv, importlib.util, os, re, urllib.request
import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), "..")
OUT_CSV = os.path.join(ROOT, "data-entry/house_district_county_louisiana.csv")
HOUSE_PAST_CSV = os.path.join(ROOT, "data-entry/house_past_results.csv")
PRES_CSV = os.path.join(ROOT, "data/county_presidential_results_2008_2024.csv")
CACHE_DIR = os.path.join(ROOT, "data-entry/la-sos-results")
# (year, election date) — the November round, then the December runoff where one was held.
ELECTIONS = [
    (2016, "2016-11-08"), (2016, "2016-12-10"),
    (2018, "2018-11-06"), (2018, "2018-12-08"),
    (2020, "2020-11-03"), (2020, "2020-12-05"),
    (2022, "2022-11-08"), (2022, "2022-12-10"),
    (2024, "2024-11-05"), (2024, "2024-12-07"),
]

spec = importlib.util.spec_from_file_location("bl", os.path.join(os.path.dirname(__file__), "build-house-district-county-results.py"))
bl = importlib.util.module_from_spec(spec)
spec.loader.exec_module(bl)
XLSX_PARTY_RE = re.compile(r"\s*\((DEM|REP|NOPTY|OTHER|IND|GRN|LBT|N/P)\)\s*$", re.I)


def workbook(date):
    y, m, d = date.split("-")
    os.makedirs(CACHE_DIR, exist_ok=True)
    local = os.path.join(CACHE_DIR, f"{y}{m}{d}.xlsx")
    if os.path.exists(local) and os.path.getsize(local) > 0:
        return local
    legacy = os.path.join(ROOT, "data-entry/la_2024_official_results.xlsx")
    if date == "2024-11-05" and os.path.exists(legacy):
        return legacy
    url = f"https://s3-us-west-2.amazonaws.com/mediaresults.sos.la.gov/HumanReadableElectionResults/{y}{m}{d}/Election+Results+({m}-{d}-{y}).xlsx"
    try:
        req = urllib.request.Request(url, headers={"User-Agent": "election-map-data-pipeline/1.0"})
        with urllib.request.urlopen(req, timeout=120) as r, open(local, "wb") as out:
            out.write(r.read())
        return local
    except Exception as e:
        print(f"  {date}: no workbook ({e})")
        return None


def parse(path, year, house_past, fips_map):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next((n for n in wb.sheetnames if "Multi-Parish" in n and "Parish)" in n), None)
    if not sheet:
        return {}
    rows = list(wb[sheet].iter_rows(values_only=True))
    starts = [i for i, r in enumerate(rows) if r and r[0] and "Congress" in str(r[0])]
    starts.append(len(rows))
    out = {}
    for si in range(len(starts) - 1):
        m = re.search(r"(\d+)(?:st|nd|rd|th) Congressional District", str(rows[starts[si]][0]))
        if not m:
            continue
        dnum = int(m.group(1))
        past = house_past.get((year, dnum))
        if not past:
            continue
        candidates = [XLSX_PARTY_RE.sub("", str(c)) for c in rows[starts[si] + 1][1:] if c]
        slots = [bl.bucket_for(c, past, {}, "LA")[0] for c in candidates]
        district_rows = {}
        for r in rows[starts[si] + 3: starts[si + 1]]:
            if not r or r[0] is None:
                break
            key = re.sub(r"\s+", " ", str(r[0])).strip().lower()
            if key not in fips_map:
                continue
            fips, name = fips_map[key]
            dem = gop = oth = 0
            for slot, v in zip(slots, r[1:1 + len(candidates)]):
                v = int(v or 0)
                if slot == "dem":
                    dem += v
                elif slot == "gop":
                    gop += v
                else:
                    oth += v
            district_rows[fips] = dict(year=year, state="LA", district=dnum, county_fips=fips, county_name=name,
                                       dem=dem, gop=gop, oth=oth, total=dem + gop + oth)
        if district_rows:
            out[dnum] = district_rows
    return out


def main():
    house_past = {}
    with open(HOUSE_PAST_CSV, newline="") as f:
        for row in csv.DictReader(f):
            if row["state_abbr"] == "LA":
                house_past[(int(row["year"]), int(row["district_name"].split("-")[1]))] = row
    fips_map = {}
    with open(PRES_CSV, newline="") as f:
        for row in csv.DictReader(f):
            if row["county_id"][:2] == "22":
                fips_map[re.sub(r"\s+", " ", row["county_name"]).strip().lower()] = (row["county_id"], row["county_name"])
    by_year = {}
    for year, date in ELECTIONS:
        path = workbook(date)
        if not path:
            continue
        parsed = parse(path, year, house_past, fips_map)
        print(f"  {date}: districts {sorted(parsed)}")
        by_year.setdefault(year, {}).update(parsed)  # a runoff overrides the November round
    rows = [r for districts in by_year.values() for d in districts.values() for r in d.values()]
    rows.sort(key=lambda r: (r["year"], r["district"], r["county_fips"]))
    with open(OUT_CSV, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["year", "state", "district", "county_fips", "county_name", "dem", "gop", "oth", "total"])
        w.writeheader()
        for r in rows:
            w.writerow(r)
    print(f"Wrote {len(rows)} rows -> {OUT_CSV}")


if __name__ == "__main__":
    main()
