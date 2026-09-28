#!/usr/bin/env python3
"""
Final-round ranked-choice results BY COUNTY for Maine's 2nd District House races decided by
RCV rounds, tabulated from the Secretary of State's own cast-vote-record files exactly as the
state did (the official summary reports only the statewide rounds; the county breakdown has to
be replicated ballot by ballot). Writes data-entry/house_district_county_maine_rcv.csv, which
scripts/build-house-district-county-results.py takes as its first-priority source, so a
ranked-choice race shows the round that decided it — never the first round.

  2018  Golden vs. Poliquin (Bond, Hoar eliminated): reuses
        scripts/fill-county-house-2018-me-cd2-rcv.py's tabulation unchanged (imported).
  2024  Golden vs. Theriault (write-ins eliminated): Nov24 CVR Export 20241115 - {1..4}.xlsx.
  2022  Golden vs. Poliquin (Bond eliminated): the SOS page publishes only the summary PDF,
        no CVR — not reproducible here, so the builder keeps that race labelled first-round.

Downloads ~25 MB from maine.gov each run. Run from project root:
  python3 scripts/build-house-district-county-maine-rcv.py
"""
import csv, importlib.util, io, os, re, urllib.request
from collections import Counter, defaultdict
import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), "..")
OUT_CSV = os.path.join(ROOT, "data-entry/house_district_county_maine_rcv.csv")
PRES_CSV = os.path.join(ROOT, "data/county_presidential_results_2008_2024.csv")

spec = importlib.util.spec_from_file_location("me18", os.path.join(os.path.dirname(__file__), "fill-county-house-2018-me-cd2-rcv.py"))
me18 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(me18)

CVR_2024 = [f"https://www.maine.gov/sos/sites/maine.gov.sos/files/inline-files/Nov24%20CVR%20Export%2020241115%20-%20{i}.xlsx" for i in (1, 2, 3, 4)]


def county_names():
    return {r["county_id"]: r["county_name"] for r in csv.DictReader(open(PRES_CSV, newline="")) if r["county_id"][:2] == "23"}


def town_of(precinct: str) -> str:
    # "Augusta W2", "Lewiston W3 P1", "Bangor D4" → the municipality the gazetteer knows
    return re.sub(r"\s+(W|Ward|D|Dist|P|Pct)\s*\d+.*$", "", str(precinct).strip(), flags=re.I).strip()


def classify_2024(cell):
    if cell is None:
        return "undervote"
    s = str(cell)
    if s in ("undervote", "overvote"):
        return s
    if s.startswith("Golden"):
        return "GOLDEN"
    if s.startswith("Theriault"):
        return "THERIAULT"
    return "OTHER"  # write-ins, eliminated in round 1


def tabulate_2024():
    """Round 2 of Golden vs. Theriault: each ballot counts for its highest-ranked continuing
    candidate; a ballot ranking only write-ins/undervotes is exhausted; an overvote ends it."""
    by_town = defaultdict(Counter)
    for url in CVR_2024:
        raw = me18.fetch_bytes(url)
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            town = town_of(row[1])
            choices = [classify_2024(c) for c in row[3:6]]
            for c in choices:
                if c in ("GOLDEN", "THERIAULT"):
                    by_town[town][c] += 1
                    break
                if c == "overvote":
                    break
                # OTHER (write-in) and undervote: look at the next ranking
    return by_town


def by_county(by_town, town_to_county):
    out = defaultdict(Counter)
    unmatched = Counter()
    for town, counts in by_town.items():
        fips = me18.match_county(town, town_to_county)
        if fips is None:
            unmatched[town] += sum(counts.values())
            continue
        out[fips].update(counts)
    return out, unmatched


def main():
    names = county_names()
    town_to_county = me18.load_town_to_county()
    rows = []

    print("2018 ME-02: tabulating the SOS cast-vote records (Golden vs. Poliquin)…", flush=True)
    r2018, un18 = by_county(me18.process_ballots(), town_to_county)
    g = sum(c["GOLDEN"] for c in r2018.values()); p = sum(c["POLIQUIN"] for c in r2018.values())
    print(f"   Golden {g} / Poliquin {p} (certified 142440 / 138931); unmatched ballots {sum(un18.values())} {dict(un18)}")
    for fips, c in sorted(r2018.items()):
        rows.append(dict(year=2018, state="ME", district=2, county_fips=fips, county_name=names.get(fips, fips),
                         dem=c["GOLDEN"], gop=c["POLIQUIN"], oth=0, total=c["GOLDEN"] + c["POLIQUIN"]))

    print("2024 ME-02: tabulating the SOS cast-vote records (Golden vs. Theriault)…", flush=True)
    r2024, un24 = by_county(tabulate_2024(), town_to_county)
    g = sum(c["GOLDEN"] for c in r2024.values()); t = sum(c["THERIAULT"] for c in r2024.values())
    print(f"   Golden {g} / Theriault {t} (certified 197151 / 194445); unmatched ballots {sum(un24.values())} {dict(list(un24.items())[:12])}")
    for fips, c in sorted(r2024.items()):
        rows.append(dict(year=2024, state="ME", district=2, county_fips=fips, county_name=names.get(fips, fips),
                         dem=c["GOLDEN"], gop=c["THERIAULT"], oth=0, total=c["GOLDEN"] + c["THERIAULT"]))

    with open(OUT_CSV, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["year", "state", "district", "county_fips", "county_name", "dem", "gop", "oth", "total"])
        w.writeheader()
        for r in rows:
            w.writerow(r)
    print(f"Wrote {len(rows)} rows -> {OUT_CSV}")


if __name__ == "__main__":
    main()
